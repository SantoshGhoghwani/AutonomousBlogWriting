import re
from pathlib import Path
from typing import Callable, Dict, Tuple

import markdown as md_lib

Meta = Dict[str, str]

CSS = """
@page { size: A4; margin: 2cm; }
body { font-family: Helvetica, Arial, sans-serif; font-size: 11pt; line-height: 1.5; color: #222; }
h1 { font-size: 22pt; margin-bottom: 4pt; }
h2 { font-size: 15pt; margin-top: 18pt; }
h3 { font-size: 12pt; }
a { color: #1a5fb4; }
code, pre { font-family: Courier; font-size: 9.5pt; background-color: #f2f2f2; }
blockquote { margin-left: 12pt; color: #555; }
table { border-collapse: collapse; }
td, th { border: 1px solid #bbb; padding: 4pt; }
"""


def split_front_matter(text: str) -> Tuple[Meta, str]:
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n(.*)$", text.strip(), re.S)
    if not m:
        return {}, text.strip()
    meta: Meta = {}
    for line in m.group(1).splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            meta[k.strip().lower()] = v.strip().strip("'\"")
    return meta, m.group(2).strip()


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "blog-post"


def _html_doc(meta: Meta, body: str) -> str:
    html_body = md_lib.markdown(body, extensions=["extra", "sane_lists", "tables", "fenced_code"])
    title = meta.get("title", "Blog Post")
    return (
        f"<html><head><meta charset='utf-8'><title>{title}</title>"
        f"<style>{CSS}</style></head><body>{html_body}</body></html>"
    )


# --- format writers -----------------------------------------------------------
def _to_md(meta: Meta, body: str, path: Path, raw: str) -> None:
    path.write_text(raw.strip() + "\n", encoding="utf-8")


def _to_html(meta: Meta, body: str, path: Path, raw: str) -> None:
    path.write_text(_html_doc(meta, body), encoding="utf-8")


def _to_pdf(meta: Meta, body: str, path: Path, raw: str) -> None:
    from xhtml2pdf import pisa

    with open(path, "wb") as f:
        status = pisa.CreatePDF(src=_html_doc(meta, body), dest=f, encoding="utf-8")
    if status.err:
        raise RuntimeError("PDF generation failed")


def _to_txt(meta: Meta, body: str, path: Path, raw: str) -> None:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", body)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"(\*\*|__|`)", "", text)
    path.write_text(text.strip() + "\n", encoding="utf-8")


_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")


def _add_runs(par, text: str) -> None:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            par.add_run(part[1:-1]).font.name = "Courier New"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def _to_docx(meta: Meta, body: str, path: Path, raw: str) -> None:
    from docx import Document

    doc = Document()
    buf: list = []

    def flush():
        if buf:
            _add_runs(doc.add_paragraph(), " ".join(buf))
            buf.clear()

    in_code = False
    for line in body.splitlines():
        if line.strip().startswith("```"):
            flush()
            in_code = not in_code
            continue
        if in_code:
            par = doc.add_paragraph()
            par.add_run(line).font.name = "Courier New"
            continue
        s = line.strip()
        if not s:
            flush()
        elif (h := re.match(r"^(#{1,6})\s+(.*)", s)):
            flush()
            doc.add_heading(h.group(2), level=min(len(h.group(1)) - 1, 4))
        elif re.match(r"^[-*+]\s+", s):
            flush()
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*+]\s+", "", s))
        elif re.match(r"^\d+[.)]\s+", s):
            flush()
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", s))
        elif s.startswith(">"):
            flush()
            _add_runs(doc.add_paragraph(style="Intense Quote"), s.lstrip("> "))
        elif re.match(r"^(-{3,}|\*{3,})$", s):
            flush()
        else:
            buf.append(s)
    flush()
    doc.save(path)


EXPORTERS: Dict[str, Tuple[str, Callable]] = {
    "markdown": (".md", _to_md),
    "md": (".md", _to_md),
    "html": (".html", _to_html),
    "pdf": (".pdf", _to_pdf),
    "docx": (".docx", _to_docx),
    "txt": (".txt", _to_txt),
}

SUPPORTED_FORMATS = sorted(set(EXPORTERS) - {"md"})


def export(markdown_text: str, fmt: str, out_dir: str, fallback_title: str = "blog post") -> Path:
    fmt = fmt.lower().strip()
    if fmt not in EXPORTERS:
        raise ValueError(f"Unsupported format '{fmt}'. Choose from: {', '.join(SUPPORTED_FORMATS)}")
    ext, writer = EXPORTERS[fmt]
    meta, body = split_front_matter(markdown_text)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"{slugify(meta.get('title', fallback_title))}{ext}"
    writer(meta, body, path, markdown_text)
    return path